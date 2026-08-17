import fitz  # PyMuPDF
import docx
import hashlib
import io
import logging
from typing import Any, Dict, List, Tuple

from config import settings

logger = logging.getLogger("airag.parser")

class DocumentParser:
    @staticmethod
    def extract_images_from_pdf(file_bytes: bytes) -> list:
        """
        抽取 PDF 中的圖片，並依序套用四道過濾：
        xref 去重 → 尺寸門檻 → 大小門檻 → 內容雜湊去重。

        跨頁重複出現的 logo／圖章與裝飾性小圖若全數送去多模態模型描述，不但會把圖片
        數量放大數倍（含大量圖片的 PDF 因此整份同步逾時），重複的描述文字還會在向量
        檢索時洗版擠掉真正相關的內容，因此重複者直接捨棄、只保留第一次出現的那張。

        單張圖片抽取失敗只記錄後跳過，不中斷其後頁面的處理。
        """
        import fitz
        images = []
        seen_xrefs = set()
        seen_hashes = set()
        total = 0
        skipped_xref = 0
        skipped_hash = 0
        skipped_size = 0
        skipped_bytes = 0
        failed = 0

        doc = None
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page_idx, page in enumerate(doc):
                image_list = page.get_images(full=True)
                for img_idx, img in enumerate(image_list):
                    total += 1
                    xref = img[0]
                    if xref in seen_xrefs:
                        skipped_xref += 1
                        continue
                    seen_xrefs.add(xref)

                    try:
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        image_ext = base_image["ext"]
                    except Exception as img_err:
                        failed += 1
                        logger.error(
                            f"[DocumentParser] 第 {page_idx + 1} 頁第 {img_idx + 1} 張圖片抽取失敗 "
                            f"(xref={xref})，已跳過: {type(img_err).__name__}: {img_err}"
                        )
                        continue

                    # 取不到寬高時視為「通過」尺寸關卡，交由大小門檻判斷；
                    # 不可預設為 0，那會讓尺寸未知的圖片被門檻靜默丟棄
                    width = base_image.get("width")
                    height = base_image.get("height")
                    if (width is not None and width < settings.IMAGE_MIN_WIDTH) or \
                       (height is not None and height < settings.IMAGE_MIN_HEIGHT):
                        skipped_size += 1
                        continue

                    if len(image_bytes) < settings.IMAGE_MIN_BYTES:
                        skipped_bytes += 1
                        continue

                    # 同一張圖以不同 xref 重複嵌入時，xref 去重擋不掉，需再比對內容雜湊
                    content_hash = hashlib.sha256(image_bytes).hexdigest()
                    if content_hash in seen_hashes:
                        skipped_hash += 1
                        continue
                    seen_hashes.add(content_hash)

                    images.append({
                        "image_bytes": image_bytes,
                        "ext": image_ext,
                        "page_number": page_idx + 1,
                        "image_index": img_idx + 1
                    })
        except Exception as e:
            logger.error(f"Error extracting images from PDF: {e}")
        finally:
            if doc is not None:
                doc.close()

        # 沒有這行統計就無法判斷門檻是否設得恰當，調整 IMAGE_MIN_* 前務必先看它
        logger.info(
            f"[DocumentParser] PDF 共抽出 {total} 張圖片：xref 重複 -{skipped_xref}、"
            f"內容重複 -{skipped_hash}、尺寸門檻 -{skipped_size}、大小門檻 -{skipped_bytes}、"
            f"抽取失敗 -{failed}，實際送描述 {len(images)} 張"
        )
        return images

    @staticmethod
    def extract_images_from_docx(file_bytes: bytes) -> list:
        """
        抽取 DOCX 中的圖片（含表格儲存格內的圖片），並儘量關聯到最近的標題階層路徑。
        """
        import docx
        from docx.oxml.ns import nsmap, qn
        import io

        # Ensure 'v' is registered in nsmap
        if 'v' not in nsmap:
            nsmap['v'] = 'urn:schemas-microsoft-com:vml'

        images = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))

            # We will walk paragraph items in order, tracking headers, to associate parent_id
            from services.word_parent_child_chunker import iter_block_items, get_heading_level, get_paragraph_text

            current_headers = {}
            img_counter = 0

            def collect_embedded_object_rids(element):
                """
                收集 <w:object>（Word 內嵌 OLE 物件，例如插入的 PDF / Excel / 簡報附件）底下的圖片 rId。
                這些圖片只是該附件在版面上顯示的「檔案圖示」或縮圖，不是文件本身的內容圖片，
                送去多模態模型只會得到無意義的描述或直接失敗，必須排除。
                """
                skip_rids = set()
                for obj in element.findall('.//' + qn('w:object')):
                    for imgdata in obj.findall('.//' + qn('v:imagedata')):
                        rid = imgdata.get(qn('r:id'))
                        if rid:
                            skip_rids.add(rid)
                    for blip in obj.findall('.//' + qn('a:blip')):
                        rid = blip.get(qn('r:embed')) or blip.get(qn('r:link'))
                        if rid:
                            skip_rids.add(rid)
                return skip_rids

            def extract_rids_from_element(element):
                rids = []
                skip_rids = collect_embedded_object_rids(element)
                # 1. Modern drawings
                blips = element.findall('.//' + qn('a:blip'))
                for blip in blips:
                    rid = blip.get(qn('r:embed')) or blip.get(qn('r:link'))
                    if rid and rid not in skip_rids:
                        rids.append(rid)
                # 2. Legacy drawings (VML)
                imagedatas = element.findall('.//' + qn('v:imagedata'))
                for imgdata in imagedatas:
                    rid = imgdata.get(qn('r:id'))
                    if rid and rid not in skip_rids:
                        rids.append(rid)
                if skip_rids:
                    logger.info(
                        f"Skipped {len(skip_rids)} embedded object icon image(s) in DOCX "
                        f"(內嵌附件檔案的圖示，非文件內容圖片)"
                    )
                return rids

            def append_images_for_rids(rids, header_path):
                nonlocal img_counter
                for rid in rids:
                    if rid in doc.part.related_parts:
                        part = doc.part.related_parts[rid]
                        if "image" in part.content_type:
                            img_counter += 1
                            ext = part.content_type.split('/')[-1]
                            if ext == "jpeg":
                                ext = "jpg"
                            images.append({
                                "image_bytes": part.blob,
                                "ext": ext,
                                "header_path": header_path.copy(),
                                "image_index": img_counter
                            })

            for item in iter_block_items(doc):
                if isinstance(item, docx.text.paragraph.Paragraph):
                    # Track headers
                    heading_level = get_heading_level(item)
                    if heading_level > 0:
                        text = get_paragraph_text(item)
                        if text:
                            # Clear lower level headers
                            for l in range(heading_level, 7):
                                current_headers.pop(f"Header {l}", None)
                            current_headers[f"Header {heading_level}"] = text

                    # Extract images in the paragraph
                    rids = extract_rids_from_element(item._element)
                    append_images_for_rids(rids, current_headers)
                elif isinstance(item, docx.table.Table):
                    # 表格儲存格內也可能內嵌圖片（常見的圖文並排排版方式），沿用目前累積到的標題階層
                    for row in item.rows:
                        for cell in row.cells:
                            for para in cell.paragraphs:
                                rids = extract_rids_from_element(para._element)
                                append_images_for_rids(rids, current_headers)

        except Exception as e:
            logger.error(f"Error extracting images from DOCX: {e}")
        return images

    @staticmethod
    def parse_pdf_pages(file_bytes: bytes) -> List[str]:
        """
        解析 PDF 檔案位元組，回傳「逐頁」純文字（索引 0 為第 1 頁）。

        切分時需要頁面邊界才能算出合理的父段落、payload 也才填得出真實頁碼；
        parse_pdf() 把所有頁面接成一整串後這些資訊就再也還原不回來了。
        """
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages = [page.get_text() for page in doc]
            doc.close()
            return pages
        except Exception as e:
            logger.error(f"Error parsing PDF: {e}")
            raise ValueError(f"PDF 解析失敗: {str(e)}")

    @staticmethod
    def parse_pdf_lines(file_bytes: bytes) -> List[List[Dict[str, Any]]]:
        """
        解析 PDF 檔案位元組，回傳「逐頁、逐行」的排版資訊：
        每行為 {"text": 文字, "size": 最大字級, "bold": 是否粗體}。

        PDF 沒有語意標記，唯一能還原標題階層的線索就是排版本身（字級、粗體）。
        get_text("text") 會把這些線索全部丟掉，只剩純文字。
        """
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            pages_lines: List[List[Dict[str, Any]]] = []
            for page in doc:
                lines: List[Dict[str, Any]] = []
                page_dict = page.get_text("dict")
                for block in page_dict.get("blocks", []):
                    # type 1 為圖片區塊，沒有文字
                    if block.get("type") != 0:
                        continue
                    for line in block.get("lines", []):
                        spans = line.get("spans", [])
                        text = "".join(span.get("text", "") for span in spans)
                        if not text.strip():
                            continue
                        max_size = max((span.get("size", 0.0) for span in spans), default=0.0)
                        # PyMuPDF span flags 的 bit 4（值 16）代表粗體
                        is_bold = any(int(span.get("flags", 0)) & 16 for span in spans)
                        lines.append({"text": text, "size": max_size, "bold": is_bold})
                pages_lines.append(lines)
            doc.close()
            return pages_lines
        except Exception as e:
            logger.error(f"Error parsing PDF layout: {e}")
            raise ValueError(f"PDF 版面解析失敗: {str(e)}")

    @staticmethod
    def lines_to_page_texts(pages_lines: List[List[Dict[str, Any]]]) -> List[str]:
        """
        把 parse_pdf_lines() 的結果還原成逐頁純文字，避免為了同時取得文字與排版而解析兩次 PDF。
        """
        return ["\n".join(line["text"] for line in lines) for lines in pages_lines]

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> Tuple[str, int]:
        """
        解析 PDF 檔案位元組，提取純文字與頁數。
        """
        pages = cls.parse_pdf_pages(file_bytes)
        return "\n".join(pages).strip(), len(pages)

    @staticmethod
    def parse_docx(file_bytes: bytes) -> Tuple[str, int]:
        """
        解析 DOCX 檔案位元組，提取純文字。
        """
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)
            # 讀取表格中的文字
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text.append(" | ".join(row_text))
            
            return "\n".join(text).strip(), 1  # Word 文件頁數通常以 1 代替或估算
        except Exception as e:
            logger.error(f"Error parsing DOCX: {e}")
            raise ValueError(f"DOCX 解析失敗: {str(e)}")

    @staticmethod
    def parse_text(file_bytes: bytes) -> Tuple[str, int]:
        """
        解析純文字/Markdown 檔案。
        """
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("big5", errors="ignore")
            except Exception as e:
                logger.error(f"Error decoding text: {e}")
                raise ValueError(f"編碼解析失敗，請確認檔案使用 UTF-8 或 Big5 編碼: {str(e)}")
        return text.strip(), 1

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> Tuple[str, int, int]:
        """
        依據副檔名自動分流解析，回傳 (純文字, 頁數, 字元數)。
        """
        ext = filename.split(".")[-1].lower()
        if ext == "pdf":
            text, pages = cls.parse_pdf(file_bytes)
        elif ext in ["docx", "doc", "dotx"]:
            from services.word_parent_child_chunker import parse_docx_to_markdown, parse_doc_to_markdown
            if ext in ["docx", "dotx"]:
                text = parse_docx_to_markdown(file_bytes)
            else:
                text = parse_doc_to_markdown(file_bytes, filename)
            pages = 1
        elif ext in ["txt", "md", "markdown", "4gl", "4fd"]:
            text, pages = cls.parse_text(file_bytes)
        else:
            raise ValueError(f"目前不支援 .{ext} 的檔案格式。支援的格式有 PDF, DOCX, DOC, DOTX, TXT, MD, 4GL, 4FD")
        
        return text, pages, len(text)

